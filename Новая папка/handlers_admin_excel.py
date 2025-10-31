# handlers_admin_excel.py
import openpyxl
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.worksheet import Worksheet
from aiogram import Router, F
from aiogram.types import CallbackQuery, Message, FSInputFile
from sqlalchemy import select
import tempfile
import os

from db.models import Session, Question, User, Event, Review, Post
from filters import IsAdminOrManager
from keyboard import admin_keyboard, manager_keyboard
from config import ADMIN_IDS

router = Router()


async def create_excel_file():
    """Создает Excel файл с данными из всех таблиц"""
    # Создаем новую книгу Excel
    wb = openpyxl.Workbook()

    # Удаляем стандартный лист
    wb.remove(wb.active)

    # Создаем листы для каждой таблицы
    await create_questions_sheet(wb)
    await create_events_sheet(wb)
    await create_reviews_sheet(wb)
    await create_posts_sheet(wb)
    await create_users_sheet(wb)

    # Сохраняем во временный файл
    with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as tmp_file:
        wb.save(tmp_file.name)
        return tmp_file.name


async def create_questions_sheet(wb: openpyxl.Workbook):
    """Создает лист с вопросами"""
    ws: Worksheet = wb.create_sheet("Вопросы")

    # Заголовки
    headers = ['ID', 'User ID', 'Username', 'First Name', 'Last Name', 'Текст вопроса',
               'Дата создания', 'Ответ', 'Дата ответа']

    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)

    # Данные
    async with Session() as session:
        result = await session.execute(
            select(Question, User)
            .join(User, Question.user_id == User.user_id)
            .order_by(Question.created_at.desc())
        )
        questions = result.all()

    for row, (question, user) in enumerate(questions, 2):
        ws.cell(row=row, column=1, value=question.id)
        ws.cell(row=row, column=2, value=question.user_id)
        ws.cell(row=row, column=3, value=user.username or '')
        ws.cell(row=row, column=4, value=user.first_name or '')
        ws.cell(row=row, column=5, value=user.last_name or '')
        ws.cell(row=row, column=6, value=question.question)
        ws.cell(row=row, column=7,
                value=question.created_at.strftime('%Y-%m-%d %H:%M:%S') if question.created_at else '')
        ws.cell(row=row, column=8, value=question.answer or '')
        ws.cell(row=row, column=9,
                value=question.answered_at.strftime('%Y-%m-%d %H:%M:%S') if question.answered_at else '')

    # Автоподбор ширины колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width


async def create_events_sheet(wb: openpyxl.Workbook):
    """Создает лист с мероприятиями"""
    ws: Worksheet = wb.create_sheet("Мероприятия")

    # Заголовки
    headers = ['ID', 'Название', 'Описание', 'Дата мероприятия', 'Ссылка на видео', 'Дата создания']

    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)

    # Данные
    async with Session() as session:
        result = await session.execute(
            select(Event).order_by(Event.event_date.desc())
        )
        events = result.scalars().all()

    for row, event in enumerate(events, 2):
        ws.cell(row=row, column=1, value=event.id)
        ws.cell(row=row, column=2, value=event.title)
        ws.cell(row=row, column=3, value=event.description)
        ws.cell(row=row, column=4, value=event.event_date.strftime('%Y-%m-%d %H:%M:%S') if event.event_date else '')
        ws.cell(row=row, column=5, value=event.video_url or '')
        ws.cell(row=row, column=6, value=event.created_at.strftime('%Y-%m-%d %H:%M:%S') if event.created_at else '')

    # Автоподбор ширины колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width


async def create_reviews_sheet(wb: openpyxl.Workbook):
    """Создает лист с отзывами"""
    ws: Worksheet = wb.create_sheet("Отзывы")

    # Заголовки
    headers = ['ID', 'User ID', 'Username', 'Event ID', 'Название мероприятия',
               'Текст отзыва', 'Дата создания']

    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)

    # Данные
    async with Session() as session:
        result = await session.execute(
            select(Review, User, Event)
            .join(User, Review.user_id == User.user_id)
            .join(Event, Review.event_id == Event.id)
            .order_by(Review.created_at.desc())
        )
        reviews = result.all()

    for row, (review, user, event) in enumerate(reviews, 2):
        ws.cell(row=row, column=1, value=review.id)
        ws.cell(row=row, column=2, value=review.user_id)
        ws.cell(row=row, column=3, value=user.username or '')
        ws.cell(row=row, column=4, value=review.event_id)
        ws.cell(row=row, column=5, value=event.title)
        ws.cell(row=row, column=6, value=review.text)
        ws.cell(row=row, column=7, value=review.created_at.strftime('%Y-%m-%d %H:%M:%S') if review.created_at else '')

    # Автоподбор ширины колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width


async def create_posts_sheet(wb: openpyxl.Workbook):
    """Создает лист с отложенными постами"""
    ws: Worksheet = wb.create_sheet("Отложенные посты")

    # Заголовки
    headers = ['ID', 'Тип', 'Медиа', 'Текст', 'Текст кнопки', 'Ссылка кнопки',
               'Время отправки', 'Статус отправки']

    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)

    # Данные
    async with Session() as session:
        result = await session.execute(
            select(Post).order_by(Post.send_at.desc())
        )
        posts = result.scalars().all()

    for row, post in enumerate(posts, 2):
        ws.cell(row=row, column=1, value=post.id)
        ws.cell(row=row, column=2, value=post.type)
        ws.cell(row=row, column=3, value=post.media or '')
        ws.cell(row=row, column=4, value=post.text or '')
        ws.cell(row=row, column=5, value=post.button_text or '')
        ws.cell(row=row, column=6, value=post.button_link or '')
        ws.cell(row=row, column=7, value=post.send_at.strftime('%Y-%m-%d %H:%M:%S') if post.send_at else '')
        ws.cell(row=row, column=8, value='Отправлено' if post.flag else 'Ожидает')

    # Автоподбор ширины колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width


async def create_users_sheet(wb: openpyxl.Workbook):
    """Создает лист с пользователями"""
    ws: Worksheet = wb.create_sheet("Пользователи")

    # Заголовки
    headers = ['User ID', 'Username', 'First Name', 'Last Name', 'Заблокирован']

    for col, header in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=header).font = Font(bold=True)

    # Данные
    async with Session() as session:
        result = await session.execute(
            select(User).order_by(User.user_id)
        )
        users = result.scalars().all()

    for row, user in enumerate(users, 2):
        ws.cell(row=row, column=1, value=user.user_id)
        ws.cell(row=row, column=2, value=user.username or '')
        ws.cell(row=row, column=3, value=user.first_name or '')
        ws.cell(row=row, column=4, value=user.last_name or '')
        ws.cell(row=row, column=5, value='Да' if user.user_is_block else 'Нет')

    # Автоподбор ширины колонок
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width


@router.callback_query(F.data == "excel", IsAdminOrManager())
async def export_to_excel(callback: CallbackQuery):
    """Экспорт данных в Excel"""
    await callback.answer("🔄 Начинаем выгрузку данных...")

    try:
        # Создаем Excel файл
        file_path = await create_excel_file()

        # Отправляем файл

        await callback.message.answer_document(
            document=FSInputFile(file_path),
            caption="📊 Выгрузка данных из бота\n\n"
                    "Файл содержит следующие листы:\n"
                    "• Вопросы - все вопросы пользователей\n"
                    "• Мероприятия - созданные мероприятия\n"
                    "• Отзывы - отзывы о мероприятиях\n"
                    "• Отложенные посты - запланированные рассылки\n"
                    "• Пользователи - информация о пользователях"
        )

        # Удаляем временный файл
        os.unlink(file_path)

        # Возвращаем в меню
        if callback.from_user.id in ADMIN_IDS:
            keyboard = admin_keyboard
        else:
            keyboard = manager_keyboard

        await callback.message.answer(
            "✅ Выгрузка завершена! Возвращаемся в меню:",
            reply_markup=keyboard
        )

    except Exception as e:
        await callback.message.answer(
            f"❌ Произошла ошибка при выгрузке: {str(e)}"
        )

        # Возвращаем в меню даже при ошибке
        if callback.from_user.id in ADMIN_IDS:
            keyboard = admin_keyboard
        else:
            keyboard = manager_keyboard

        await callback.message.answer(
            "Возвращаемся в меню:",
            reply_markup=keyboard
        )
